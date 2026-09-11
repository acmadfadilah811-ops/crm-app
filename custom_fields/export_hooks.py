"""
Runtime hooks so custom fields appear in Horilla's Select Columns to Export
modal and are written into exported files.

Horilla's export catalog only knows about model columns. This module patches
those call sites from the custom_fields app so we do not edit Horilla sources.
"""

import csv
import logging
from io import BytesIO, StringIO

from django.db.models.query import QuerySet
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from custom_fields.detail_hooks import custom_field_selector_items
from custom_fields.list_hooks import attach_custom_field_values_to_objects
from horilla.contrib.core.utils import sanitize_export_value

logger = logging.getLogger(__name__)

_PATCHED = False


def _cf_property(name):
    return property(lambda obj, key=name: obj.__dict__.get(key, ""))


def _install_export_properties(model, extras):
    """Expose ``cf_*`` names as properties so Horilla's export catalog includes them."""
    installed = []
    for _label, name in extras:
        if not hasattr(model, name):
            setattr(model, name, _cf_property(name))
            installed.append(name)
    old_labels = getattr(model, "PROPERTY_LABELS", None)
    labels = dict(old_labels or {})
    for label, name in extras:
        labels[name] = label
    model.PROPERTY_LABELS = labels
    return installed, old_labels


def _uninstall_export_properties(model, installed, old_labels):
    for name in installed:
        if hasattr(model, name):
            delattr(model, name)
    if old_labels is None:
        if hasattr(model, "PROPERTY_LABELS"):
            delattr(model, "PROPERTY_LABELS")
    else:
        model.PROPERTY_LABELS = old_labels


def _cell_values(obj, extra_pairs):
    values = []
    for _label, name in extra_pairs:
        raw = obj.__dict__.get(name, "")
        values.append(sanitize_export_value("" if raw is None else str(raw)))
    return values


def inject_custom_fields_into_export_modules(modules):
    """Add custom fields to Settings → Export Data column pickers."""
    from horilla.apps import apps

    for module in modules or []:
        app_label = module.get("app_label")
        model_name = module.get("name")
        if not app_label or not model_name:
            continue
        try:
            model = apps.get_model(app_label, model_name)
        except LookupError:
            continue
        extras = custom_field_selector_items(model)
        if not extras:
            continue
        fields = list(module.get("fields") or [])
        existing = {item.get("name") for item in fields}
        for label, name in extras:
            if name not in existing:
                fields.append({"name": name, "label": label})
                existing.add(name)
        module["fields"] = fields
    return modules


def _custom_fields_only_export(view, model, export_format, objects, extra_pairs):
    """Build a csv/xlsx buffer that contains only custom-field columns."""
    headers = [str(label) for label, _name in extra_pairs]
    rows = [_cell_values(obj, extra_pairs) for obj in objects]
    filename = view.get_export_filename(model, export_format)
    if export_format == "csv":
        buffer = StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(rows)
        return filename, BytesIO(buffer.getvalue().encode("utf-8"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center")
    header_fill = PatternFill(
        start_color="eafb5b", end_color="eafb5b", fill_type="solid"
    )
    for cell in sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.fill = header_fill
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return filename, buffer


def _append_custom_field_columns(data, export_format, objects, extra_pairs):
    """Append custom-field columns onto an already-built csv/xlsx export buffer."""
    extra_headers = [str(label) for label, _name in extra_pairs]
    extra_values = [_cell_values(obj, extra_pairs) for obj in objects]
    raw = data.getvalue()
    if export_format == "csv":
        text = raw.decode("utf-8") if isinstance(raw, (bytes, bytearray)) else raw
        reader = csv.reader(StringIO(text))
        rows = list(reader)
        if not rows:
            rows = [[]]
        rows[0] = list(rows[0]) + extra_headers
        for index, extra_row in enumerate(extra_values):
            if index + 1 < len(rows):
                rows[index + 1] = list(rows[index + 1]) + extra_row
            else:
                rows.append(extra_row)
        buffer = StringIO()
        csv.writer(buffer).writerows(rows)
        return BytesIO(buffer.getvalue().encode("utf-8"))

    workbook = load_workbook(BytesIO(raw))
    sheet = workbook.active
    start_col = sheet.max_column + 1
    for offset, header in enumerate(extra_headers):
        sheet.cell(row=1, column=start_col + offset, value=header)
    for row_index, extra_row in enumerate(extra_values, start=2):
        for offset, value in enumerate(extra_row):
            sheet.cell(row=row_index, column=start_col + offset, value=value)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def install_export_patches():
    """Monkey-patch Horilla export helpers without editing their files."""
    global _PATCHED
    if _PATCHED:
        return

    from horilla.contrib.core.views.export_data import ExportView, get_export_cell_value
    from horilla.contrib.generics.views.toolkit.bulk_export import (
        HorillaBulkExportMixin,
    )

    original_handle_export = HorillaBulkExportMixin.handle_export
    original_get_available_models = ExportView.get_available_models
    original_export_model_data = ExportView.export_model_data
    original_get_export_cell_value = get_export_cell_value

    def patched_handle_export(self, record_ids, columns, export_format):
        model = getattr(self, "model", None)
        extras = custom_field_selector_items(model) if model is not None else []
        if not extras:
            return original_handle_export(self, record_ids, columns, export_format)

        installed, old_labels = _install_export_properties(model, extras)
        orig_iter = QuerySet.__iter__

        def attaching_iter(qs):
            iterator = orig_iter(qs)
            if getattr(qs, "model", None) is not model:
                return iterator
            items = list(iterator)
            try:
                attach_custom_field_values_to_objects(model, items, extras=extras)
            except Exception:
                logger.exception("custom_fields: could not attach export values")
            return iter(items)

        QuerySet.__iter__ = attaching_iter
        try:
            return original_handle_export(self, record_ids, columns, export_format)
        finally:
            QuerySet.__iter__ = orig_iter
            _uninstall_export_properties(model, installed, old_labels)

    def patched_get_available_models(self):
        modules = original_get_available_models(self)
        try:
            inject_custom_fields_into_export_modules(modules)
        except Exception:
            logger.exception("custom_fields: could not inject export columns")
        return modules

    def patched_export_model_data(
        self, model, export_format, queryset=None, selected_fields=None
    ):
        extras = custom_field_selector_items(model)
        extra_names = {name for _label, name in extras}
        extra_pairs = [
            (label, name)
            for label, name in extras
            if selected_fields is None or name in selected_fields
        ]
        if not extra_pairs:
            return original_export_model_data(
                self, model, export_format, queryset, selected_fields
            )

        objects = list(model.objects.all() if queryset is None else queryset)
        attach_custom_field_values_to_objects(model, objects, extras=extras)

        model_selected = None
        if selected_fields is not None:
            model_selected = [
                name for name in selected_fields if name not in extra_names
            ]
            if not model_selected:
                return _custom_fields_only_export(
                    self, model, export_format, objects, extra_pairs
                )

        filename, data = original_export_model_data(
            self, model, export_format, objects, model_selected
        )
        if extra_pairs and export_format in ("csv", "xlsx"):
            data = _append_custom_field_columns(
                data, export_format, objects, extra_pairs
            )
        return filename, data

    def patched_get_export_cell_value(obj, field_name, field, user):
        if str(field_name).startswith("cf_"):
            value = obj.__dict__.get(field_name, "")
            return "" if value is None else str(value)
        return original_get_export_cell_value(obj, field_name, field, user)

    HorillaBulkExportMixin.handle_export = patched_handle_export
    ExportView.get_available_models = patched_get_available_models
    ExportView.export_model_data = patched_export_model_data

    import horilla.contrib.core.views.export_data as export_data_mod

    export_data_mod.get_export_cell_value = patched_get_export_cell_value
    _PATCHED = True
